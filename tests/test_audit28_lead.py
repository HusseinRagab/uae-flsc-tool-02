"""Audit 28.10.2026, lead-session fixes: mixed most-restrictive mode (owner ruling R-F).

R-F: one Table 6.5 emergency lighting package that meets every row (#4.2, p.419) and one
Table 8.13 alarm protocol per building (extending R7)."""
import os
import sys

sys.path.insert(0, os.path.join(os.path.dirname(__file__), ".."))

from engine import evaluate  # noqa: E402
from flsc_schema import Building  # noqa: E402

_MIX = dict(mixed_occupancy_mode="most_restrictive", floors_above_grade=4,
            ground_floor_bua_m2=2000, gross_floor_area_m2=8000,
            secondary_storage_hazard="OH1", sum_gf_bua_m2=2000)


def _block(kw, chap, prefix):
    r = evaluate(Building(**{**_MIX, **kw}))
    ch = next(c for c in r.chapters if c.chapter_code == chap)
    return [it for bl in ch.blocks if bl.title.startswith(prefix) for it in bl.items]


def _el_types(items):
    return [it.system for it in items if it.status == "required" and (
        it.system.startswith(("Monitored Self-Contained EL", "Central Battery",
                              "Stand-alone Self-Contained EL")))]


def test_RF_el_mall_plus_lowrise_keeps_monitored_only():
    # Table 6.5 item 3 (mall: Central Battery OR Monitored) + item 6 (low rise: Monitored)
    items = _block(dict(occupancy="assembly_a", secondary_occupancies=["mall_covered"],
                        height_m=12), "EL", "EL - System")
    types = _el_types(items)
    assert types == ["Monitored Self-Contained EL (Section 3.3)"]
    line = next(it for it in items if it.system == types[0])
    assert "items 3 and 6" in line.code_ref and "R-F" in line.detail


def test_RF_el_highrise_plus_fuel_drops_stand_alone():
    # item 2 (Central Battery OR Monitored) + item 13 fuel station (stand-alone): #4.2
    items = _block(dict(occupancy="residential", secondary_occupancies=["motor_fuel_dispensing"],
                        height_m=30), "EL", "EL - System")
    types = _el_types(items)
    assert len(types) == 1 and types[0].startswith("Central Battery")
    assert any("#4.2" in (it.system or "") for it in items)


def test_RF_el_single_occupancy_unchanged():
    items = _block(dict(occupancy="mall_covered", secondary_occupancies=[],
                        mixed_occupancy_mode="none", height_m=12), "EL", "EL - System")
    assert len(_el_types(items)) == 1


def _protocols(items):
    return [(it.source_rule, it.system) for it in items if it.status == "required" and (
        "Phased Evacuation Alarm" in it.system or it.system.startswith(
            ("Positive Alarm Sequence", "General Fire Alarm")))]


def test_RF_fa_residential_plus_hotel_one_protocol():
    # Table 8.13 item 2 (5-floor phased, PAS 15/180/240 s) vs item 18 hotel (no evacuation
    # on general alarm, PAS 120 s): the hotel row governs (R7), the height protocol is named.
    items = _block(dict(occupancy="residential", secondary_occupancies=["hotel_a"],
                        height_m=30), "FA", "FA - Detection")
    p = _protocols(items)
    assert {r for r, _ in p} == {"fa_hotel"}
    assert any("5-Floor Phased Evacuation Alarm" in (it.detail or "") for it in items)


def test_RF_fa_residential_plus_storage_one_protocol():
    # items 3-4 PAS 180 s vs items 19-20 (vii) general alarm = evacuation alarm
    items = _block(dict(occupancy="residential", secondary_occupancies=["storage_industrial"],
                        height_m=10), "FA", "FA - Detection")
    p = _protocols(items)
    assert {r for r, _ in p} == {"fa_storage_industrial"}


def test_RF_fa_single_highrise_keeps_height_protocol():
    items = _block(dict(occupancy="residential", secondary_occupancies=[],
                        mixed_occupancy_mode="none", height_m=30), "FA", "FA - Detection")
    names = [s for _, s in _protocols(items)]
    assert "5-Floor Phased Evacuation Alarm" in names


def test_RF_evc_mixed_one_package():
    # residential + mall at 30 m: Table 7.3 items 2 and 3 print one EVC and one two-way line
    items = _block(dict(occupancy="residential", secondary_occupancies=["mall_covered"],
                        height_m=30), "EVC", "EVC - Voice")
    req = [it for it in items if it.status == "required"]
    assert len(req) == 2
    assert all("items 2 and 3" in it.code_ref for it in req)


def test_RB_fa_general_lines_conditional_for_tank_farm_tank_detection_required():
    r = evaluate(Building(occupancy="storage_industrial", height_m=10, floors_above_grade=1,
                          ground_floor_bua_m2=2000, gross_floor_area_m2=2000, sum_gf_bua_m2=2000,
                          hazard_class="HH", is_aboveground_tank_storage=True))
    fa = next(c for c in r.chapters if c.chapter_code == "FA")
    gen = [it for bl in fa.blocks if bl.title.startswith("FA - General") for it in bl.items]
    assert gen and all(it.status != "required" for it in gen)
    eq = [it for bl in fa.blocks if bl.title.startswith("FA - Equipment") for it in bl.items]
    assert any(it.status == "required" and "Table 8.15 item 3" in (it.code_ref or "") for it in eq)


def _fp_systems(**kw):
    r = evaluate(Building(**kw))
    fp = next(c for c in r.chapters if c.chapter_code == "FP")
    return [it.system for bl in fp.blocks for it in bl.items]


def test_F9A03_plot_blank_not_evaluated_tall_midrise_villa():
    # Tables 9.18 / 9.19 / 9.20.a / 9.23.B split on the plot area (p.696-700, p.717):
    # a blank plot names the large-plot row as NOT EVALUATED (standing blank rule).
    for kw in (dict(occupancy="residential", height_m=60, floors_above_grade=18),
               dict(occupancy="residential", height_m=120, floors_above_grade=35),
               dict(occupancy="assembly_a", height_m=20, floors_above_grade=5),
               dict(occupancy="villa_commercial", height_m=8, floors_above_grade=2)):
        sy = _fp_systems(ground_floor_bua_m2=1000, gross_floor_area_m2=5000, **kw)
        assert any("Plot Area" in s and "NOT EVALUATED" in s for s in sy), kw
        sy = _fp_systems(ground_floor_bua_m2=1000, gross_floor_area_m2=5000, plot_area_m2=3000, **kw)
        assert not any("Plot Area" in s and "NOT EVALUATED" in s for s in sy), kw


def test_Q4_fa_facp_line_follows_storage_tenancy():
    def fa(t):
        r = evaluate(Building(occupancy="storage_industrial", height_m=8, ground_floor_bua_m2=2000,
                              gross_floor_area_m2=2000, sum_gf_bua_m2=2000, plot_area_m2=4000,
                              hazard_class="OH1", storage_tenancy=t))
        c = next(c for c in r.chapters if c.chapter_code == "FA")
        return {it.source_rule for bl in c.blocks for it in bl.items}
    assert "fa_storage_single_tenant_facp" in fa("single")
    assert "fa_storage_multi_tenant_facp" in fa("multi")
    assert not {"fa_storage_single_tenant_facp", "fa_storage_multi_tenant_facp"} & fa(None)


def test_UX03_not_evaluated_lines_have_their_own_status():
    # A missing input is never counted as a required system (audit UX-03).
    r = evaluate(Building(occupancy="residential", height_m=18, floors_above_grade=5,
                          ground_floor_bua_m2=0, gross_floor_area_m2=0))
    ne = [it for c in r.chapters for bl in c.blocks for it in bl.items
          if "NOT EVALUATED" in it.system]
    assert ne and all(it.status == "not_evaluated" for it in ne)


def test_UX03_exports_say_incomplete_and_list_not_evaluated_lines():
    import io
    import docx
    import export
    r = evaluate(Building(occupancy="residential", height_m=18, floors_above_grade=5,
                          ground_floor_bua_m2=0, gross_floor_area_m2=0))
    for kind in ("compact", "detailed", "executive"):
        d = docx.Document(io.BytesIO(export.report_to_docx_bytes(r, kind)))
        text = "\n".join(p.text for p in d.paragraphs)
        assert "REPORT INCOMPLETE" in text, kind
        assert "NOT EVALUATED" in text, kind
    ok = evaluate(Building(occupancy="residential", height_m=18, floors_above_grade=5,
                           ground_floor_bua_m2=800, gross_floor_area_m2=4000, plot_area_m2=2000))
    d = docx.Document(io.BytesIO(export.report_to_docx_bytes(ok, "compact")))
    assert "REPORT INCOMPLETE" not in "\n".join(p.text for p in d.paragraphs)
    assert export.report_to_pdf_bytes(r, "compact")[:4] == b"%PDF"


def test_bugbot_villa_not_mandatory_line_dropped_beside_facility_el():
    # PR #38 review: villa + fuel station with a basement printed the villa "not mandatory"
    # line beside the fuel station's required stand-alone line (one kind only).
    r = evaluate(Building(occupancy="villa_private", secondary_occupancies=["motor_fuel_dispensing"],
                          mixed_occupancy_mode="most_restrictive", height_m=8, floors_above_grade=2,
                          floors_below_grade=1, basement_bua_m2=300, ground_floor_bua_m2=400,
                          gross_floor_area_m2=800, plot_area_m2=1000))
    el = next(c for c in r.chapters if c.chapter_code == "EL")
    its = [it for bl in el.blocks if bl.title.startswith("EL - System") for it in bl.items]
    assert any(it.status == "required" for it in its)
    assert not any(it.status == "not_required" and it.system.startswith("Emergency Lighting (")
                   for it in its)
    # Two villas: only the basement stand-alone line is required; the villa line stays.
    r2 = evaluate(Building(occupancy="villa_private", secondary_occupancies=["villa_commercial"],
                           mixed_occupancy_mode="most_restrictive", height_m=8, floors_above_grade=2,
                           floors_below_grade=1, basement_bua_m2=300, ground_floor_bua_m2=400,
                           gross_floor_area_m2=800, plot_area_m2=1000))
    el2 = next(c for c in r2.chapters if c.chapter_code == "EL")
    assert any(it.status == "not_required" and it.system.startswith("Emergency Lighting (")
               for bl in el2.blocks for it in bl.items)


def test_bugbot_parking_duty_promoted_when_host_has_none():
    # PR #38 review, ruling Q-3: a host whose own table prints no pump / tank takes the
    # attached parking's duty into its Ch 9 chapter; the parking chapter keeps none.
    from engine import _merge_attached_parking_duties
    from flsc_schema import ChapterReport, Requirement, SectionBlock
    host = [ChapterReport(chapter_code="FP", chapter_title="FP", blocks=[SectionBlock(
        title="FP - Branch Requirements (Ch 9)", items=[Requirement(system="Hose Reel System")])])]
    park = [ChapterReport(chapter_code="FP", chapter_title="FP", blocks=[SectionBlock(
        title="FP - Branch Requirements (Ch 9)", items=[
            Requirement(system="Fire Pump", spec="750 gpm @ 6.9 bar", duty="pump"),
            Requirement(system="Fire Water Tank", spec="60 minutes duration", duty="tank")])])]
    h, p = _merge_attached_parking_duties(host, park)
    duties = [it.duty for bl in h[0].blocks for it in bl.items if it.duty]
    assert sorted(duties) == ["pump", "tank"]
    assert not [it for bl in p[0].blocks for it in bl.items if it.duty]
    assert all("Q-3" in it.detail for bl in h[0].blocks for it in bl.items if it.duty)


def test_bugbot_exports_do_not_count_not_evaluated_as_required():
    # PR #38 review: compact chapter bars and the executive parking row said "N required"
    # with NOT EVALUATED lines inside N (audit UX-03).
    import io
    import docx
    import export
    r = evaluate(Building(occupancy="residential", height_m=18, floors_above_grade=5,
                          ground_floor_bua_m2=0, gross_floor_area_m2=0))
    fp = next(c for c in r.chapters if c.chapter_code == "FP")
    items = export._required_systems(fp)
    n_req = sum(1 for it in items if it.status == "required")
    n_ne = sum(1 for it in items if it.status == "not_evaluated")
    assert n_ne and export._count_text(items) == f"{n_req} required, {n_ne} not evaluated"
    d = docx.Document(io.BytesIO(export.report_to_docx_bytes(r, "compact")))
    cells = " ".join(c.text for t in d.tables for row in t.rows for c in row.cells)
    assert f"{n_req} required, {n_ne} not evaluated" in cells


# ------------------------------------------------ independent verifier, 28.10.2026
def test_verifier_conditional_not_evaluated_lines_get_the_status():
    # Ch 10 builds some NOT EVALUATED lines as conditional (e.g. basement depth blank).
    r = evaluate(Building(occupancy="residential", height_m=30, floors_above_grade=8,
                          floors_below_grade=2, basement_bua_m2=1000, ground_floor_bua_m2=1000,
                          gross_floor_area_m2=8000, plot_area_m2=3000))
    ne = [it for c in r.chapters for bl in c.blocks for it in bl.items
          if "NOT EVALUATED" in it.system]
    assert ne and all(it.status == "not_evaluated" for it in ne)


def test_verifier_gated_9_29_9_30_lines_not_evaluated_when_row_is():
    r = evaluate(Building(occupancy="healthcare_a", height_m=8, floors_above_grade=2,
                          ground_floor_bua_m2=2000, gross_floor_area_m2=4000, plot_area_m2=4000,
                          tent_marquee_kind="small_detached"))
    fp = next(c for c in r.chapters if c.chapter_code == "FP")
    its = [it for bl in fp.blocks for it in bl.items]
    assert not [it for it in its if it.status == "not_required" and (
        "Table 9.29" in (it.code_ref or "") or "Table 9.30" in (it.code_ref or ""))]
    assert any(it.status == "not_evaluated" and it.system.startswith("Lift Machine Room")
               for it in its)


def test_verifier_unranked_pumps_name_each_other():
    r = evaluate(Building(occupancy="storage_industrial", height_m=8, floors_above_grade=1,
                          ground_floor_bua_m2=2000, gross_floor_area_m2=2000, sum_gf_bua_m2=2000,
                          plot_area_m2=4000, hazard_class="OH2", storage_tenancy="multi",
                          is_cold_storage=True))
    fp = next(c for c in r.chapters if c.chapter_code == "FP")
    pumps = [it for bl in fp.blocks for it in bl.items if it.duty == "pump" and it.status == "required"]
    if len(pumps) > 1:
        assert all("shall meet each of them" in (it.detail or "") for it in pumps)


def test_verifier_RD_highrise_hotel_kitchens_required():
    r = evaluate(Building(occupancy="hotel_a", height_m=40, floors_above_grade=12,
                          ground_floor_bua_m2=1500, gross_floor_area_m2=18000, plot_area_m2=5000))
    sc = next(c for c in r.chapters if c.chapter_code == "SC")
    rules = {it.source_rule: it.status for bl in sc.blocks for it in bl.items}
    assert rules.get("sc_aux_kitchen_16iii") == "required"
    assert rules.get("sc_aux_kitchen_16iii_residential_highrise") == "conditional"


def test_verifier_secondary_storage_tenancy_reaches_the_storage_run():
    kw = dict(occupancy="residential", secondary_occupancies=["storage_industrial"],
              mixed_occupancy_mode="most_restrictive", height_m=12, floors_above_grade=3,
              ground_floor_bua_m2=2000, gross_floor_area_m2=6000, plot_area_m2=4000,
              secondary_storage_hazard="OH1", sum_gf_bua_m2=2000)
    def ne(**x):
        r = evaluate(Building(**kw, **x))
        return [it.system for c in r.chapters for bl in c.blocks for it in bl.items
                if "Tenancy" in it.system and it.status == "not_evaluated"]
    assert ne()                       # blank: named, never a pass
    assert not ne(storage_tenancy="multi")


def test_verifier_app_asks_tenancy_for_secondary_storage():
    src = open(os.path.join(os.path.dirname(__file__), "..", "app.py"), encoding="utf-8").read()
    i = src.index('key="secondary_storage_hazard_input"')
    assert "_storage_tenancy_input()" in src[i:i + 900]
    assert 'occupancy == "storage_industrial" or _storage_secondary)\n        and not st.session_state.get("storage_tenancy_input")' in src


def test_verifier_no_itm_pointer_for_a_not_evaluated_system():
    r = evaluate(Building(occupancy="healthcare_a", height_m=8, floors_above_grade=2,
                          ground_floor_bua_m2=2000, gross_floor_area_m2=4000, plot_area_m2=4000,
                          tent_marquee_kind="small_detached"))
    fp = next(c for c in r.chapters if c.chapter_code == "FP")
    rules = {it.source_rule for bl in fp.blocks for it in bl.items}
    assert "itm_sprinkler" not in rules and "itm_clean_agent" not in rules


def test_verifier_recheck_outlet_area_blank_gated_lines_not_evaluated():
    # Business 12 m, outlet ticked, outlet area blank: Table 9.21.e row 2 (over 900 m2, p.711)
    # would require sprinklers, so the Table 9.29 / 9.30 lines are NOT EVALUATED, not a pass.
    r = evaluate(Building(occupancy="business", business_group="B", height_m=12,
                          floors_above_grade=3, ground_floor_bua_m2=800, gross_floor_area_m2=2400,
                          plot_area_m2=2000, has_attached_commercial_outlet=True))
    fp = next(c for c in r.chapters if c.chapter_code == "FP")
    its = [it for bl in fp.blocks for it in bl.items]
    assert not [it for it in its if it.status == "not_required" and (
        "Table 9.29 row D" in (it.code_ref or "") or "Table 9.30.M" in (it.code_ref or ""))]
    src = open(os.path.join(os.path.dirname(__file__), "..", "app.py"), encoding="utf-8").read()
    assert 'has_attached_commercial_outlet")\n        and float(st.session_state.get("commercial_outlet_gf_bua_m2_input")' in src


def test_verifier_recheck_parking_run_keeps_the_choice_with_a_carried_tick():
    r = evaluate(Building(occupancy="mall_covered", height_m=18, floors_above_grade=3,
                          floors_below_grade=2, basement_bua_m2=25000, depth_below_grade_m=7,
                          ground_floor_bua_m2=25000, gross_floor_area_m2=75000, plot_area_m2=50000,
                          has_theater_cinema=True, theater_cinema_kind="cinema",
                          has_parking_area=True, parking_area_m2=25000, parking_height_m=0,
                          parking_basement_levels=2, parking_depth_m=7))
    ne = [it.system for c in list(r.chapters) + list(r.attached_parking_chapters)
          for bl in c.blocks for it in bl.items if it.status == "not_evaluated"]
    assert not ne, ne


def test_verifier_recheck_first_load_sample_is_a_complete_report():
    from streamlit.testing.v1 import AppTest
    at = AppTest.from_file(os.path.join(os.path.dirname(__file__), "..", "app.py"),
                           default_timeout=180)
    at.run()
    assert not at.exception and not at.error
    assert not [w for w in at.warning if "Report incomplete" in w.value]
